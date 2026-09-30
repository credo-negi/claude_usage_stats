#!/usr/bin/env python3
"""標準ライブラリの unittest だけで動くテスト（実ログ・~/.claude・ネットワークには触れない）。

    python tests/test_usage.py          # または  python -m unittest discover -s tests
"""
import contextlib
import io
import json
import os
import sys
import tempfile
import unittest
from datetime import date

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import claude_usage  # noqa: E402
from cusage import aggregate, hook, logs, pricing, report_html, store  # noqa: E402

SESSION = "sess-0001"
CWD = "/work/app"


def assistant(ts, mid, rid, model="claude-sonnet-5", inp=10, out=100, cw5=0, cw1=0, cr=0, side=False,
              speed="standard", cwd=CWD, session=SESSION, cache_creation=True, uuid=None):
    usage = {"input_tokens": inp, "output_tokens": out, "cache_read_input_tokens": cr,
             "cache_creation_input_tokens": cw5 + cw1, "speed": speed}
    if cache_creation:
        usage["cache_creation"] = {"ephemeral_5m_input_tokens": cw5, "ephemeral_1h_input_tokens": cw1}
    return {"type": "assistant", "timestamp": ts, "sessionId": session, "cwd": cwd, "isSidechain": side,
            "requestId": rid, "uuid": uuid or f"u-{mid}",
            "message": {"id": mid, "model": model, "role": "assistant", "content": [], "usage": usage}}


def human(ts, text, uuid, session=SESSION, cwd=CWD, origin=True, **extra):
    d = {"type": "user", "timestamp": ts, "sessionId": session, "cwd": cwd, "uuid": uuid,
         "message": {"role": "user", "content": [{"type": "text", "text": text}]}}
    if origin is True:
        d["origin"] = {"kind": "human"}
    elif isinstance(origin, dict):
        d["origin"] = origin
    d.update(extra)
    return d


def tool_result(ts, uuid):
    return {"type": "user", "timestamp": ts, "sessionId": SESSION, "cwd": CWD, "uuid": uuid,
            "message": {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "t", "content": "ok"}]}}


def read_text(path):
    with open(path, "r", encoding="utf-8") as fh:
        return fh.read()


def read_json(path):
    return json.loads(read_text(path))


def write_jsonl(path, events, raw_tail=""):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        for e in events:
            fh.write(json.dumps(e, ensure_ascii=False, separators=(",", ":")) + "\n")
        fh.write(raw_tail)


def make_config(tmp):
    """~/.claude 相当のディレクトリを合成する。"""
    cfg = os.path.join(tmp, "claude")
    main = os.path.join(cfg, "projects", "-work-app", SESSION + ".jsonl")
    sub = os.path.join(cfg, "projects", "-work-app", SESSION, "subagents", "agent-a1.jsonl")
    events = [
        # ターン1: 2026-09-01 09:00 JST (= 00:00Z)
        human("2026-09-01T00:00:00.000Z", "最初の質問 <ide_opened_file>x.py</ide_opened_file>", "p1"),
        assistant("2026-09-01T00:00:05.000Z", "m1", "r1", inp=10, out=100, cw5=1000, cw1=2000, cr=5000),
        # ストリーミングで同じ応答が 2 行（重複排除の対象。後の行の output が完成形）
        assistant("2026-09-01T00:00:06.000Z", "m2", "r2", inp=5, out=10, cr=100),
        assistant("2026-09-01T00:00:07.000Z", "m2", "r2", inp=5, out=50, cr=100),
        tool_result("2026-09-01T00:00:08.000Z", "tr1"),
        {"type": "assistant", "timestamp": "2026-09-01T00:00:09.000Z", "sessionId": SESSION,
         "message": {"id": "m3", "model": "<synthetic>", "usage": {"input_tokens": 1, "output_tokens": 1}}},
        # 人間ではない user 行（ターンを増やしてはいけない）
        human("2026-09-01T00:00:10.000Z", "<local-command-stdout>ok</local-command-stdout>", "n1", origin=False),
        human("2026-09-01T00:00:11.000Z", "meta", "n2", isMeta=True),
        human("2026-09-01T00:00:12.000Z", "bg done", "n3", origin={"kind": "task-notification"}),
        # ターン2: JST では 9/2 に日付が変わる（2026-09-01T15:30Z = 9/2 00:30 JST）
        human("2026-09-01T15:30:00.000Z", "<command-name>/review</command-name><command-args>PR 12</command-args>", "p2", origin=False),
        assistant("2026-09-01T15:30:05.000Z", "m4", "r4", model="claude-opus-5-5", inp=100, out=1000, speed="fast"),
        assistant("2026-09-01T15:30:06.000Z", "m5", "r5", model="claude-unknown-9", inp=1, out=1),
        # cache_creation の内訳が無い古い形式（全部 5 分扱い）
        assistant("2026-09-01T15:30:07.000Z", "m6", "r6", model="claude-haiku-4-5-20251001", inp=0, out=0, cw5=400, cache_creation=False),
    ]
    write_jsonl(main, events, raw_tail='{"type":"assistant","truncated')       # 書き込み途中の末尾行
    # サブエージェントは別ファイル。ターン2の最中に動いた
    write_jsonl(sub, [assistant("2026-09-01T15:31:00.000Z", "s1", "sr1", model="claude-haiku-4-5", inp=200, out=300,
                                side=True, cwd="")])
    return cfg


class LogsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.cfg = make_config(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_parse_dedup_and_filters(self):
        path = os.path.join(self.cfg, "projects", "-work-app", SESSION + ".jsonl")
        calls, prompts = logs.parse_file(path, "-work-app")
        by = {c.key: c for c in calls}
        self.assertEqual(set(by), {"m1:r1", "m2:r2", "m4:r4", "m5:r5", "m6:r6"})   # <synthetic> は除外
        self.assertEqual(by["m2:r2"].out, 50)                                         # 重複は出力が大きい方
        self.assertEqual((by["m1:r1"].cw5, by["m1:r1"].cw1), (1000, 2000))
        self.assertEqual(by["m6:r6"].cw5, 400)                                        # 内訳なし → 5m
        self.assertEqual(by["m4:r4"].speed, "fast")
        # 人間のプロンプトは 2 件だけ（tool_result / local-command / isMeta / task-notification は除外）
        self.assertEqual([p.uuid for p in prompts], ["p1", "p2"])
        self.assertEqual(prompts[0].preview, "最初の質問")                              # IDE タグは除去
        self.assertEqual(prompts[1].preview, "/review PR 12")

    def test_find_transcripts_includes_subagents(self):
        found = [os.path.basename(p) for p, _d in logs.find_transcripts([self.cfg])]
        self.assertEqual(sorted(found), ["agent-a1.jsonl", SESSION + ".jsonl"])

    def test_default_dirs_env(self):
        self.assertEqual(logs.default_claude_dirs({"CLAUDE_CONFIG_DIR": self.cfg + ",/nonexistent"}), [self.cfg])


class PricingTest(unittest.TestCase):
    def setUp(self):
        self.p = pricing.Pricing.load()

    def call(self, **kw):
        base = dict(key="k", session="s", ts="2026-09-01T00:00:00Z", model="claude-sonnet-5", inp=0, out=0, cw5=0,
                    cw1=0, cr=0, speed="standard", cwd="", side=0, pdir="")
        base.update(kw)
        return logs.Call(**base)

    def test_resolution(self):
        r = self.p.resolve
        self.assertEqual(r("claude-haiku-4-5-20251001").key, "claude-haiku-4-5")
        self.assertEqual(r("anthropic.claude-opus-5-5").key, "claude-opus-5-5")
        self.assertEqual(r("us.anthropic.claude-sonnet-4-5-20250929-v1:0").key, "claude-sonnet-4-5")
        self.assertEqual(r("claude-opus-4-1-20250805").key, "claude-opus-4-1")
        self.assertEqual(r("claude-opus-4-20250514").key, "claude-opus-4")
        self.assertEqual(r("claude-sonnet-5-5").key, "claude-sonnet-5-5")
        self.assertIsNone(r("gpt-5"))

    def test_cost_sonnet(self):
        # sonnet-5: 入力 $2 / 出力 $10 / 書込5m $2.5 / 書込1h $4 / 読込 $0.2 （100万トークンあたり）
        c = self.call(inp=1_000_000, out=1_000_000, cw5=1_000_000, cw1=1_000_000, cr=1_000_000)
        listed, net = self.p.cost(c)
        self.assertAlmostEqual(listed, 2 + 10 + 2.5 + 4 + 0.2)
        self.assertAlmostEqual(net, listed)

    def test_fast_and_discount(self):
        c = self.call(model="claude-opus-5-5", inp=1_000_000, out=1_000_000, speed="fast")
        self.assertAlmostEqual(self.p.cost(c)[0], (4 + 20) * 2)
        p = pricing.Pricing.load(discount_override=0.25)
        listed, net = p.cost(self.call(inp=1_000_000))
        self.assertAlmostEqual(net, listed * 0.75)

    def test_unknown_model_is_zero(self):
        self.assertEqual(self.p.cost(self.call(model="mystery", inp=10**6)), (0.0, 0.0))

    def test_discount_must_be_fraction(self):
        for bad in (15, 1.0, -0.1):
            with self.assertRaises(ValueError):
                pricing.Pricing([], discount_override=bad)
        with self.assertRaises(ValueError) as cm:
            pricing.Pricing([{"model": "claude-x", "input_per_mtok": "1", "output_per_mtok": "2", "discount_rate": "15", "_line": 7}])
        self.assertIn("7 行目", str(cm.exception))                                     # どの行が悪いか分かる

    def test_csv_loading(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "p.csv")
            with open(path, "w", encoding="utf-8-sig", newline="") as fh:             # Excel の BOM 付き UTF-8
                fh.write("# コメント行\n"
                         "model,input_per_mtok,output_per_mtok,cache_read_per_mtok,fast_multiplier,discount_rate,note\n"
                         "*,,,,,0.1,既定割引\n"
                         "claude-a,10,50,,,,\"カンマ, 入り\"\n"
                         "claude-b,2,10,0.5,3,0.2,\n")
            p = pricing.Pricing.load(path)
            a, b = p.resolve("claude-a"), p.resolve("claude-b-20260101")
            self.assertEqual((a.cache_write_5m, a.cache_write_1h, a.cache_read), (12.5, 20.0, 1.0))   # 空欄は導出
            self.assertEqual((a.discount_rate, a.note), (0.1, "カンマ, 入り"))                       # * 行が既定
            self.assertEqual((b.discount_rate, b.cache_read, b.fast_multiplier), (0.2, 0.5, 3.0))
            self.assertEqual(pricing.Pricing.load(path, discount_override=0.3).resolve("claude-b").discount_rate, 0.3)
            with open(path, "w", encoding="utf-8") as fh:
                fh.write("model,input_per_mtok,output_per_mtok\nclaude-a,abc,5\n")
            with self.assertRaises(ValueError) as cm:
                pricing.Pricing.load(path)
            self.assertIn("input_per_mtok", str(cm.exception))


class AggregateTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.cfg = make_config(self.tmp.name)
        self.db = os.path.join(self.tmp.name, "u.db")
        s = store.Store(self.db)
        store.sync(s, [self.cfg])
        self.calls, self.prompts = s.load_calls(), s.load_prompts()
        s.close()
        self.pricing = pricing.Pricing.load()

    def tearDown(self):
        self.tmp.cleanup()

    def report(self, tz="Asia/Tokyo", **kw):
        return aggregate.build_report(self.calls, self.prompts, self.pricing, tz, **kw)

    def test_turns_include_subagent_from_other_file(self):
        rep = self.report()
        self.assertEqual(len(rep.turns), 2)
        t1, t2 = rep.turns
        self.assertEqual((t1.no, t1.bucket.calls), (1, 2))
        self.assertEqual((t2.no, t2.bucket.calls, t2.bucket.sub_calls), (2, 4, 1))    # サブエージェントは turn 2 に合算
        self.assertEqual(t2.project, CWD)                                              # サブエージェントの cwd 欠落でも親セッションのもの
        self.assertEqual(rep.total.turns, 2)

    def test_turn_duration(self):
        rep = self.report()
        t1, t2 = rep.turns
        self.assertAlmostEqual(t1.duration, 7.0)       # 00:00:00 の送信 → 最後の応答 m2(00:00:07)
        self.assertAlmostEqual(t2.duration, 7.0)       # 後から動いたサブエージェント(15:31:00)では延びない
        self.assertEqual(rep.duration_stats(), (14.0, 2))
        orphan = aggregate.build_report(self.calls, [], self.pricing, "UTC")
        self.assertIsNone(orphan.turns[0].duration)    # プロンプト不明のターンは算出しない
        self.assertEqual(orphan.duration_stats(), (0, 0))
        self.assertEqual([t[5] for t in report_html.build_payload(rep)["turns"]], [7.0, 7.0])

    def test_token_and_cost_totals(self):
        rep = self.report()
        m = {k: v for k, v in rep.models.items()}
        # sonnet-5: m1 (10 in,100 out,1000 cw5,2000 cw1,5000 cr) + m2 (5,50,cr 100)
        exp_sonnet = (15 * 2 + 150 * 10 + 1000 * 2.5 + 2000 * 4 + 5100 * 0.2) / 1e6
        self.assertAlmostEqual(m["claude-sonnet-5"].cost, exp_sonnet)
        self.assertIn("claude-opus-5-5 (fast)", m)                                     # Fast は別行
        self.assertAlmostEqual(m["claude-opus-5-5 (fast)"].cost, (100 * 4 + 1000 * 20) * 2 / 1e6)
        self.assertEqual(rep.unpriced_models, {"claude-unknown-9": 1})
        self.assertAlmostEqual(rep.total.cost, sum(b.cost for b in rep.models.values()))
        self.assertAlmostEqual(rep.total.cost, sum(b.cost for b in rep.daily.values()))
        self.assertAlmostEqual(rep.total.cost, sum(t.bucket.cost for t in rep.turns))
        self.assertAlmostEqual(rep.total.cost, sum(f.cost for f in rep.facts))
        self.assertAlmostEqual(rep.total.cost, sum(c for _t, c in rep.kind_costs.values()))

    def test_timezone_day_boundary(self):
        jst = self.report("Asia/Tokyo")
        utc = self.report("UTC")
        self.assertEqual(list(jst.daily), ["2026-09-01", "2026-09-02"])     # 15:30Z は JST では翌日
        self.assertEqual(list(utc.daily), ["2026-09-01"])
        self.assertAlmostEqual(jst.total.cost, utc.total.cost)

    def test_filters(self):
        self.assertEqual(list(self.report(filters=aggregate.Filters(since=date(2026, 9, 2))).daily), ["2026-09-02"])
        r = self.report(filters=aggregate.Filters(model="opus"))
        self.assertEqual(r.total.calls, 1)
        self.assertEqual(self.report(filters=aggregate.Filters(project="nomatch")).total.calls, 0)

    def test_prompt_privacy_default(self):
        self.assertTrue(all(t.prompt == "" for t in self.report().turns))
        self.assertEqual(self.report(prompt_chars=5).turns[0].prompt, "最初の質問")

    def test_orphan_calls_without_prompt(self):
        rep = aggregate.build_report(self.calls, [], self.pricing, "UTC")
        self.assertEqual(len(rep.turns), 1)                                            # プロンプト不明 → セッション毎の 1 ターン
        self.assertEqual(rep.turns[0].no, 0)


class StoreTest(unittest.TestCase):
    def test_incremental_and_retention(self):
        with tempfile.TemporaryDirectory() as tmp:
            cfg = make_config(tmp)
            s = store.Store(os.path.join(tmp, "u.db"))
            st1 = store.sync(s, [cfg])
            self.assertEqual(st1.files_parsed, 2)
            st2 = store.sync(s, [cfg])
            self.assertEqual(st2.files_parsed, 0)                                       # 変更なしなら再解析しない
            n = len(s.load_calls())
            # 元ログが削除されても蓄積済みの履歴は残る（Claude Code の自動削除対策）
            os.remove(os.path.join(cfg, "projects", "-work-app", SESSION + ".jsonl"))
            store.sync(s, [cfg])
            self.assertEqual(len(s.load_calls()), n)
            # ファイルが伸びたときだけ再解析し、二重計上しない
            sub = os.path.join(cfg, "projects", "-work-app", SESSION, "subagents", "agent-a1.jsonl")
            with open(sub, "a", encoding="utf-8") as fh:
                fh.write(json.dumps(assistant("2026-09-01T15:32:00.000Z", "s2", "sr2", side=True)) + "\n")
            st3 = store.sync(s, [cfg])
            self.assertEqual(st3.files_parsed, 1)
            self.assertEqual(len(s.load_calls()), n + 1)
            s.close()


class HtmlTest(unittest.TestCase):
    def test_embedded_payload_is_safe_and_consistent(self):
        with tempfile.TemporaryDirectory() as tmp:
            cfg = make_config(tmp)
            # プロンプトに </script> を混ぜても HTML が壊れない
            path = os.path.join(cfg, "projects", "-work-app", "evil.jsonl")
            write_jsonl(path, [human("2026-09-03T00:00:00.000Z", "</script><script>alert(1)</script>", "e1", session="evil"),
                               assistant("2026-09-03T00:00:01.000Z", "e-m", "e-r", session="evil", inp=1, out=1)])
            s = store.Store(os.path.join(tmp, "u.db"))
            store.sync(s, [cfg])
            rep = aggregate.build_report(s.load_calls(), s.load_prompts(), pricing.Pricing.load(), "UTC", prompt_chars=100)
            s.close()
            html = report_html.render_html(rep)
            self.assertEqual(html.count("</script>"), 2)                                # データ用と本体用の 2 つだけ
            start = html.index('<script id="data" type="application/json">') + len('<script id="data" type="application/json">')
            data = json.loads(html[start:html.index("</script>", start)])
            self.assertAlmostEqual(sum(f[10] for f in data["facts"]), rep.total.cost, places=4)
            self.assertEqual(sum(f[8] for f in data["facts"]), rep.total.calls)
            self.assertIn("</script><script>alert(1)</script>", [t[4] for t in data["turns"]])
            self.assertEqual(len(data["turns"]), len(rep.turns))


class XlsxTest(unittest.TestCase):
    def test_workbook(self):
        try:
            import openpyxl
            from cusage.report_xlsx import write_xlsx
        except ImportError:
            self.skipTest("XlsxWriter / openpyxl が無い")
        with tempfile.TemporaryDirectory() as tmp:
            cfg = make_config(tmp)
            s = store.Store(os.path.join(tmp, "u.db"))
            store.sync(s, [cfg])
            rep = aggregate.build_report(s.load_calls(), s.load_prompts(), pricing.Pricing.load(), "Asia/Tokyo")
            s.close()
            out = write_xlsx(rep, os.path.join(tmp, "r.xlsx"))
            wb = openpyxl.load_workbook(out)
            self.assertEqual(wb.sheetnames, ["サマリー", "日次", "月次", "プロジェクト別", "モデル別", "クロス集計", "ターン明細", "単価表"])
            rows = list(wb["日次"].iter_rows(min_row=2, values_only=True))
            self.assertEqual([r[0] for r in rows[:2]], ["2026-09-01", "2026-09-02"])
            self.assertAlmostEqual(sum(r[-1] for r in rows[:2]), rep.total.cost)
            self.assertEqual(wb["ターン明細"].max_row, 1 + len(rep.turns))

    def test_empty_report(self):
        try:
            from cusage.report_xlsx import write_xlsx
        except ImportError:
            self.skipTest("XlsxWriter が無い")
        rep = aggregate.build_report([], [], pricing.Pricing.load(), "UTC")
        with tempfile.TemporaryDirectory() as tmp:
            write_xlsx(rep, os.path.join(tmp, "e.xlsx"))
            self.assertIn('"facts":[]', report_html.render_html(rep))              # 空でも例外なく生成できる
            self.assertEqual(rep.total.calls, 0)


class HookTest(unittest.TestCase):
    def test_install_uninstall_preserves_others(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "settings.json")
            other = {"type": "command", "command": "/opt/other-hook.sh"}
            with open(path, "w", encoding="utf-8") as fh:
                json.dump({"model": "sonnet", "hooks": {"Stop": [{"hooks": [other]}], "PreToolUse": [{"hooks": []}]}}, fh)
            cmd = hook.build_command("/x/claude_usage.py", ["--formats", "html", "--prompt-chars", "0"])
            msg, backup = hook.install(path, cmd)
            self.assertIn("追加", msg)
            self.assertTrue(os.path.exists(backup))
            msg2, _ = hook.install(path, cmd)                                           # 冪等
            self.assertIn("更新", msg2)
            data = read_json(path)
            cmds = [h["command"] for g in data["hooks"]["Stop"] for h in g["hooks"]]
            self.assertEqual(len(cmds), 2)
            self.assertIn("/opt/other-hook.sh", cmds)
            self.assertEqual(data["model"], "sonnet")
            self.assertTrue(any(c.endswith("hook --formats html --prompt-chars 0") for c in cmds))
            n, _ = hook.uninstall(path)
            self.assertEqual(n, 1)
            data = read_json(path)
            self.assertEqual(data["hooks"]["Stop"], [{"hooks": [other]}])
            self.assertIn("PreToolUse", data["hooks"])

    def test_install_creates_missing_file_and_rejects_invalid_json(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, ".claude", "settings.json")
            hook.install(path, "python claude_usage.py hook")
            self.assertTrue(os.path.exists(path))
            with open(path, "w") as fh:
                fh.write("{ // comment\n}")
            with self.assertRaises(ValueError):
                hook.install(path, "python claude_usage.py hook")

    def test_backup_never_overwrites_previous_backup(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "settings.json")
            with open(path, "w", encoding="utf-8") as fh:
                fh.write('{"model": "original"}')
            _, b1 = hook.install(path, "python claude_usage.py hook")
            _, b2 = hook.install(path, "python claude_usage.py hook --formats html")   # 同じ秒でも別ファイル
            self.assertNotEqual(b1, b2)
            self.assertEqual(read_json(b1), {"model": "original"})

    def test_lock(self):
        with tempfile.TemporaryDirectory() as tmp:
            lock = os.path.join(tmp, ".lock")
            with hook.single_instance(lock) as a:
                with hook.single_instance(lock) as b:
                    self.assertTrue(a)
                    self.assertFalse(b)
            with hook.single_instance(lock) as c:
                self.assertTrue(c)


class CliTest(unittest.TestCase):
    def run_cli(self, *argv, env=None):
        buf = io.StringIO()
        old = dict(os.environ)
        os.environ.update(env or {})
        try:
            with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(io.StringIO()):
                rc = claude_usage.main(list(argv))
        finally:
            os.environ.clear()
            os.environ.update(old)
        return rc, buf.getvalue()

    def test_report_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp:
            cfg = make_config(tmp)
            out = os.path.join(tmp, "out")
            rc, text = self.run_cli("--claude-dir", cfg, "--timezone", "Asia/Tokyo", "--out-dir", out,
                                    env={"CLAUDE_USAGE_HOME": os.path.join(tmp, "home")})
            self.assertEqual(rc, 0)
            self.assertTrue(os.path.getsize(os.path.join(out, "claude-usage-report.html")) > 1000)
            self.assertTrue(os.path.exists(os.path.join(out, "claude-usage-report.xlsx")))
            self.assertIn("claude-unknown-9", text)                                     # 単価未設定の警告

    def test_hook_never_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            home = os.path.join(tmp, "home")
            rc, _ = self.run_cli("hook", "--claude-dir", os.path.join(tmp, "missing"), "--pricing",
                                 os.path.join(tmp, "nope.csv"), env={"CLAUDE_USAGE_HOME": home})
            self.assertEqual(rc, 0)                                                     # 失敗してもセッションを妨げない
            self.assertIn("Traceback", read_text(os.path.join(home, "hook.log")))

    def test_install_hook_rejects_bad_options(self):
        with tempfile.TemporaryDirectory() as tmp:
            settings = os.path.join(tmp, "settings.json")
            for bad in (["--formats", "bogus"], ["--discount", "15"], ["--timezone", "Mars/Base"]):
                rc, _ = self.run_cli("install-hook", "--settings", settings, *bad)
                self.assertEqual(rc, 1, bad)
            self.assertFalse(os.path.exists(settings))                                   # 不正な設定は書き込まない
            rc, _ = self.run_cli("install-hook", "--settings", settings, "--formats", "html")
            self.assertEqual(rc, 0)
            rc, _ = self.run_cli("uninstall-hook", "--settings", settings)
            self.assertEqual(rc, 0)

    def test_bad_discount_reports_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            cfg = make_config(tmp)
            rc, _ = self.run_cli("--claude-dir", cfg, "--discount", "15", "--formats", "none",
                                 env={"CLAUDE_USAGE_HOME": os.path.join(tmp, "home")})
            self.assertEqual(rc, 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
